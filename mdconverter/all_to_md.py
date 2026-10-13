"""Converte varios formatos de arquivo para Markdown (alem de PDF).

Cobre os formatos mais comuns do dia a dia: Word (.docx), Excel/CSV
(.xlsx, .csv), HTML (.html/.htm), texto puro (.txt) e RTF (.rtf). PDF
continua sendo tratado por `mdconverter.pdf_to_md` (usado por PDFtoMD/,
allPDFtoMD/ e tambem por ALLtoMD/, que aceita qualquer um destes formatos
misturados na mesma pasta).
"""

from __future__ import annotations

import csv
from pathlib import Path

SUPPORTED_SUFFIXES = (".pdf", ".docx", ".xlsx", ".csv", ".html", ".htm", ".txt", ".rtf")


def _rows_to_markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    header, *body = rows
    width = len(header)
    out = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in body:
        row = list(row) + [""] * (width - len(row))
        out.append("| " + " | ".join(cell.replace("\n", " ") for cell in row[:width]) + " |")
    return "\n".join(out)


def _docx_to_markdown(path: Path) -> str:
    import docx  # python-docx

    document = docx.Document(str(path))
    blocks: list[str] = []
    list_buffer: list[str] = []

    def flush_list() -> None:
        if list_buffer:
            blocks.append("\n".join(list_buffer))
            list_buffer.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading") or style.startswith("title"):
            flush_list()
            digits = "".join(ch for ch in style if ch.isdigit())
            level = min(6, max(1, int(digits))) if digits else 1
            blocks.append(f"{'#' * level} {text}")
        elif style.startswith("list bullet"):
            list_buffer.append(f"- {text}")
        elif style.startswith("list number"):
            list_buffer.append(f"1. {text}")
        else:
            flush_list()
            blocks.append(text)
    flush_list()

    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append(_rows_to_markdown_table(rows))

    return "\n\n".join(b for b in blocks if b.strip())


def _xlsx_to_markdown(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), data_only=True)
    sections = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if all(cell is None for cell in row):
                continue
            rows.append(["" if cell is None else str(cell) for cell in row])
        if not rows:
            continue
        table_md = _rows_to_markdown_table(rows)
        heading = f"## {sheet.title}" if len(workbook.worksheets) > 1 else f"# {path.stem}"
        sections.append(f"{heading}\n\n{table_md}")
    return "\n\n".join(sections)


def _csv_to_markdown(path: Path) -> str:
    with path.open(newline="", encoding="utf-8-sig") as f:
        rows = [row for row in csv.reader(f) if row]
    return _rows_to_markdown_table(rows)


def _html_to_markdown(path: Path) -> str:
    from markdownify import markdownify

    html_text = path.read_text(encoding="utf-8", errors="replace")
    return markdownify(html_text, heading_style="ATX", bullets="-").strip()


def _txt_to_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").strip()


def _rtf_to_markdown(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text

    rtf_source = path.read_text(encoding="utf-8", errors="replace")
    return rtf_to_text(rtf_source).strip()


_CONVERTERS = {
    ".docx": _docx_to_markdown,
    ".xlsx": _xlsx_to_markdown,
    ".csv": _csv_to_markdown,
    ".html": _html_to_markdown,
    ".htm": _html_to_markdown,
    ".txt": _txt_to_markdown,
    ".rtf": _rtf_to_markdown,
}


def convert_any_to_md(input_path: Path, output_path: Path) -> str:
    """Converte .docx / .xlsx / .csv / .html / .htm / .txt / .rtf em Markdown.

    PDFs nao passam por aqui -- use `mdconverter.pdf_to_md.convert_pdf_file`.
    Retorna a extensao de origem (para log), e grava o resultado em
    `output_path`.
    """
    suffix = input_path.suffix.lower()
    converter = _CONVERTERS.get(suffix)
    if converter is None:
        raise ValueError(f"formato nao suportado: {suffix}")

    markdown_text = converter(input_path).strip()
    if not markdown_text.startswith("#"):
        markdown_text = f"# {input_path.stem}\n\n{markdown_text}"
    markdown_text += "\n"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(markdown_text, encoding="utf-8")
    return suffix.lstrip(".")
